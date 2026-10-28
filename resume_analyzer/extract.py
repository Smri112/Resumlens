"""extract.py - read a resume file and pull structured information out of it."""
import csv
import io
import re
from pathlib import Path

import pdfplumber
from docx import Document

SKILLS_FILE = Path(__file__).parent / "data" / "skills.csv"

SECTION_KEYWORDS = {
    "Summary": ["summary", "objective", "profile", "about me"],
    "Education": ["education", "academic", "qualification"],
    "Skills": ["skills", "technical skills", "technologies", "tech stack"],
    "Projects": ["projects", "project work", "academic projects"],
    "Experience": ["experience", "internship", "work history", "employment"],
    "Certifications": ["certification", "certificate", "courses", "achievements", "awards"],
}


def read_resume(uploaded_file):
    """Return (text, page_count) from an uploaded PDF or DOCX file."""
    name = uploaded_file.name.lower()
    data = uploaded_file.getvalue()
    if name.endswith(".pdf"):
        pages = []
        with pdfplumber.open(io.BytesIO(data)) as pdf:
            for page in pdf.pages:
                pages.append(page.extract_text() or "")
        return "\n".join(pages), len(pdf.pages)
    if name.endswith(".docx"):
        doc = Document(io.BytesIO(data))
        parts = [p.text for p in doc.paragraphs]
        for table in doc.tables:
            for row in table.rows:
                parts.extend(cell.text for cell in row.cells)
        text = "\n".join(parts)
        return text, max(1, round(len(text.split()) / 500))  # rough page estimate
    raise ValueError("Unsupported file type. Please upload a PDF or DOCX.")


def load_skills():
    """Return list of (skill, category) from the skills CSV."""
    with open(SKILLS_FILE, newline="", encoding="utf-8") as f:
        return [(r["skill"], r["category"]) for r in csv.DictReader(f)]


def find_skills(text):
    """Return the set of known skills (original casing) found in the text."""
    text_l = text.lower()
    found = set()
    for skill, _ in load_skills():
        pattern = r"(?<![a-z0-9+#.])" + re.escape(skill.lower()) + r"(?![a-z0-9+#])"
        if re.search(pattern, text_l):
            found.add(skill)
    return found


def extract_email(text):
    m = re.search(r"[\w.+-]+@[\w-]+\.[\w.-]+", text)
    return m.group(0) if m else None


def extract_phone(text):
    m = re.search(r"(?:\+?\d{1,3}[\s-]?)?(?:\(?\d{3,5}\)?[\s-]?)?\d{3,5}[\s-]?\d{4,5}", text)
    if m:
        digits = re.sub(r"\D", "", m.group(0))
        if 10 <= len(digits) <= 13:
            return m.group(0).strip()
    return None


def extract_name(text):
    """Guess the name: first short line made only of letters."""
    for line in text.splitlines()[:6]:
        line = line.strip()
        words = line.split()
        if 2 <= len(words) <= 4 and all(w.replace(".", "").isalpha() for w in words):
            return line.title()
    return None


def extract_cgpa(text):
    """Find a CGPA/GPA on a 10-point scale, e.g. 'CGPA: 8.4' or '8.4/10'. Returns float or None."""
    m = re.search(r"(?:cgpa|sgpa|gpa)\s*[:\-]?\s*(\d{1,2}(?:\.\d{1,2})?)", text, re.I) \
        or re.search(r"\b(\d{1,2}(?:\.\d{1,2})?)\s*/\s*10\b", text)
    if m:
        value = float(m.group(1))
        if 0 < value <= 10:
            return value
    return None


def detect_sections(text):
    """Return {section: True/False} depending on whether a heading is present."""
    lines = [l.strip().lower() for l in text.splitlines() if l.strip()]
    result = {}
    for section, keys in SECTION_KEYWORDS.items():
        result[section] = any(
            len(line) < 40 and any(k in line for k in keys) for line in lines
        )
    return result


def extract_info(text):
    low = text.lower()
    return {
        "name": extract_name(text),
        "email": extract_email(text),
        "phone": extract_phone(text),
        "cgpa": extract_cgpa(text),
        "linkedin": "linkedin.com" in low,
        "github": "github.com" in low,
        "skills": find_skills(text),
        "sections": detect_sections(text),
    }
