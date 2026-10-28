"""make_samples.py - creates sample_data/ with fake resumes (PDF + DOCX) and a job description.
Run once:  python make_samples.py     (the files are already included in the project)
"""
from pathlib import Path

from docx import Document
from reportlab.lib.pagesizes import A4
from reportlab.lib.styles import ParagraphStyle, getSampleStyleSheet
from reportlab.lib.units import cm
from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer

OUT = Path(__file__).parent / "sample_data"
OUT.mkdir(exist_ok=True)

JD = """Machine Learning Intern

We are looking for a student with strong Python and Machine Learning fundamentals.
You will build and test models, clean data and present results to the team.

Requirements:
- Python, Pandas, NumPy and Scikit-learn
- Knowledge of Machine Learning and Statistics
- SQL for working with data
- Git and GitHub for version control
- Experience with TensorFlow or Deep Learning is a plus
- Good communication and teamwork
"""

RESUMES = {
    "aarav_mehta_strong_ml.pdf": """Aarav Mehta
aarav.mehta@example.com | +91 98765 11111 | github.com/aaravm | linkedin.com/in/aaravm
SUMMARY
Second year CSE student focused on machine learning and data analysis.
EDUCATION
B.Tech Computer Science and Engineering, 2nd year, CGPA: 8.7
SKILLS
Python, SQL, Pandas, NumPy, Scikit-learn, Machine Learning, Statistics, TensorFlow, Git, GitHub, Matplotlib, Communication, Teamwork
PROJECTS
- Built a house price predictor using Python and Scikit-learn, reducing error by 18%
- Developed a sentiment analysis tool with NLP and NLTK used by 150+ students
- Designed a sales dashboard using Pandas and Matplotlib
- Deployed a Streamlit app and presented results to 40 classmates
CERTIFICATIONS
- Machine Learning course, Coursera
- Winner, college hackathon 2025""",

    "diya_nair_web_developer.pdf": """Diya Nair
diya.nair@example.com | +91 98765 22222 | github.com/diyan
SUMMARY
Second year CSE student who enjoys building web applications.
EDUCATION
B.Tech Computer Science and Engineering, 2nd year, CGPA: 8.1
SKILLS
HTML, CSS, JavaScript, React, Node.js, MongoDB, Git, GitHub, Bootstrap, REST API, Python, SQL
PROJECTS
- Built a portfolio website with React and Bootstrap, visited by 300+ users
- Developed a task manager using Node.js, Express and MongoDB
- Created a REST API with authentication for a library system
- Collaborated with a team of 4 using Git and GitHub
EXPERIENCE
- Web development intern, 2 months: improved page load time by 25%""",

    "rohan_gupta_average.pdf": """Rohan Gupta
rohan.gupta@example.com | 9876533333
EDUCATION
B.Tech CSE, 2nd year, CGPA: 6.9
SKILLS
Python, Java, C Programming, HTML, Excel
PROJECTS
- Made a calculator in Python
- Worked on a student management system in Java
- Responsible for the college website team
ACHIEVEMENTS
- Hard worker and team player""",

    "kabir_singh_weak.pdf": """Kabir Singh
Kabir Singh, CSE student
I am a hard worker and a team player, looking for a job.
Interests: cricket, music, travelling, movies, various tasks, etc.
Languages known: English, Hindi
References available on request.""",
}


def make_pdf(name, text):
    styles = getSampleStyleSheet()
    body = ParagraphStyle("b", parent=styles["BodyText"], fontSize=10.5, leading=15)
    head = ParagraphStyle("h", parent=styles["Heading2"], fontSize=12, spaceBefore=10, spaceAfter=4)
    doc = SimpleDocTemplate(str(OUT / name), pagesize=A4, leftMargin=2 * cm, rightMargin=2 * cm,
                            topMargin=2 * cm, bottomMargin=2 * cm)
    story = []
    for i, line in enumerate(text.split("\n")):
        line = line.replace("&", "&amp;")
        if i == 0:
            story.append(Paragraph(f"<b><font size=18>{line}</font></b>", body))
        elif line.isupper() and len(line) < 25:
            story.append(Paragraph(line, head))
        else:
            story.append(Paragraph(line, body))
        story.append(Spacer(1, 1))
    doc.build(story)


for fname, content in RESUMES.items():
    make_pdf(fname, content)

# one DOCX sample to test Word support
d = Document()
d.add_heading("Meera Iyer", 0)
d.add_paragraph("meera.iyer@example.com | +91 98765 44444 | github.com/meerai")
for head, items in [
    ("Education", ["B.Tech CSE, 2nd year, CGPA: 8.4"]),
    ("Skills", ["Python, SQL, Pandas, NumPy, Data Analysis, Statistics, Power BI, Excel, Git"]),
    ("Projects", ["Built a COVID data dashboard using Python and Power BI, used by 60+ students",
                  "Analyzed 10,000 records with Pandas and improved report time by 30%"]),
]:
    d.add_heading(head, 1)
    for it in items:
        d.add_paragraph(it)
d.save(OUT / "meera_iyer_data_analyst.docx")

(OUT / "job_description_ml_intern.txt").write_text(JD, encoding="utf-8")
print("created", len(list(OUT.iterdir())), "files in sample_data/")
